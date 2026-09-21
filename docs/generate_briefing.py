"""Generate the Nagaad Municipal Risk Dashboard briefing (PPTX) and process guide (DOCX)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRgb
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches as PptInches, Pt as PptPt

OUT_DIR = Path(__file__).resolve().parent
LOGO = OUT_DIR / "nagaad-logo.png"
NAVY = PptRgb(0x1E, 0x4D, 0x7B)
SKY = PptRgb(0xE0, 0xF2, 0xFE)
WHITE = PptRgb(0xFF, 0xFF, 0xFF)
SLATE = PptRgb(0x33, 0x41, 0x55)
MUTED = PptRgb(0x64, 0x74, 0x8B)
GREEN = PptRgb(0x05, 0x96, 0x69)
DOC_NAVY = RGBColor(0x1E, 0x4D, 0x7B)
TODAY = date(2026, 9, 21).strftime("%d %B %Y")


def _set_run(run, text, size=18, bold=False, color=SLATE, font="Calibri"):
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_bar(slide, prs):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), prs.slide_width, PptInches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        PptInches(0),
        prs.slide_height - PptInches(0.42),
        prs.slide_width,
        PptInches(0.42),
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    tf = foot.text_frame
    tf.margin_left = PptInches(0.4)
    p = tf.paragraphs[0]
    p.text = ""
    run = p.add_run()
    _set_run(
        run,
        "Nagaad  ·  Somalia Urban Resilience Project Phase II  ·  Municipal Flood Risk Dashboard",
        size=11,
        color=WHITE,
    )


def _add_logo(slide, left, top, size):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), PptInches(left), PptInches(top), PptInches(size), PptInches(size))


def _blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bar(slide, prs)
    _add_logo(slide, 12.42, 0.22, 0.72)
    return slide


def _title(slide, text, top=0.28, size=28):
    box = slide.shapes.add_textbox(PptInches(0.4), PptInches(top), PptInches(11.8), PptInches(0.7))
    p = box.text_frame.paragraphs[0]
    p.clear()
    run = p.add_run()
    _set_run(run, text, size=size, bold=True, color=NAVY)
    return box


def _bullets(slide, lines, left=0.5, top=1.15, width=12.2, height=5.4, size=18):
    box = slide.shapes.add_textbox(PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 1 if line.startswith("  ") else 0
        p.space_after = PptPt(8)
        run = p.add_run()
        _set_run(run, line.strip(), size=size - (2 if line.startswith("  ") else 0), color=SLATE)
    return box


def _card(slide, x, y, w, h, title, body, fill=SKY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = PptInches(0.16)
    tf.margin_right = PptInches(0.12)
    tf.margin_top = PptInches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    _set_run(run, title, size=14, bold=True, color=NAVY)
    p2 = tf.add_paragraph()
    run2 = p2.add_run()
    _set_run(run2, body, size=13, color=SLATE)


def _table(slide, rows, left, top, width, height, col_w=None):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    table = table_shape.table
    if col_w:
        for i, w in enumerate(col_w):
            table.columns[i].width = PptInches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            _set_run(run, val, size=11, bold=(r == 0), color=WHITE if r == 0 else SLATE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (SKY if r % 2 else WHITE)
    return table


def build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    # 1 Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, PptInches(5.9), prs.slide_width, PptInches(1.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = PptRgb(0x16, 0x3A, 0x5C)
    accent.line.fill.background()
    badge = s.shapes.add_shape(
        MSO_SHAPE.OVAL, PptInches(0.55), PptInches(0.35), PptInches(1.35), PptInches(1.35)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = WHITE
    badge.line.fill.background()
    _add_logo(s, 0.62, 0.42, 1.22)
    t = s.shapes.add_textbox(PptInches(0.7), PptInches(1.9), PptInches(12), PptInches(1.2))
    p = t.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, "Nagaad Municipal Risk Dashboard", size=36, bold=True, color=WHITE)
    st = s.shapes.add_textbox(PptInches(0.7), PptInches(3.1), PptInches(12), PptInches(1.4))
    p = st.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(
        run,
        "Flood screening for Beledweyne  ·  what it does and how it is built",
        size=20,
        color=SKY,
    )
    ft = s.shapes.add_textbox(PptInches(0.7), PptInches(6.2), PptInches(12), PptInches(0.9))
    p = ft.text_frame.paragraphs[0]
    run = p.add_run()
    _set_run(run, f"Nagaad Project  ·  Hiraan  ·  {TODAY}", size=16, color=WHITE)
    p2 = ft.text_frame.add_paragraph()
    run = p2.add_run()
    _set_run(run, "Indicative screening — not a full probabilistic risk model", size=13, color=SKY)

    # 2 Agenda
    s = _blank(prs)
    _title(s, "What this briefing covers")
    _bullets(
        s,
        [
            "Why a separate municipal flood dashboard is needed",
            "What municipal staff and Nagaad see on the map and KPIs",
            "How OSM roads and Nagaad investments are analysed separately",
            "Technical architecture: web app, map tiles, Docker ETL",
            "How mixed GIS formats (GeoPackage, KML/KMZ) are converted",
            "How the dashboard is embedded in the Nagaad MIS",
            "Current Beledweyne results, limits, and next layers",
            "Improvements and data the analysis team should collect so this becomes a municipal decision tool",
        ],
    )

    # 3 Purpose
    s = _blank(prs)
    _title(s, "What the dashboard does")
    _bullets(
        s,
        [
            "Gives Beledweyne a city-scoped flood screening view for municipal decision-making",
            "Shows where historical Shabelle flood extents overlap people, buildings, and roads",
            "Separates the city-wide OSM road network from Nagaad / SURP-II invested and design-ready roads",
            "Highlights IDP and host sites with the largest populations in the flood extent",
            "Lives as its own free/open-source app, then is embedded in Nagaad MIS (iframe)",
            "This version is flood only — conflict and other hazards are out of scope",
        ],
        size=17,
    )

    # 4 Scope
    s = _blank(prs)
    _title(s, "Scope of this version")
    _card(s, 0.4, 1.15, 4.0, 2.3, "In scope", "Beledweyne (Hiraan)\nHistorical Shabelle flood\nBuildings, IDP sites, roads\nNagaad roads and PIU office")
    _card(s, 4.65, 1.15, 4.0, 2.3, "Intentionally later", "Hospitals, schools, markets\n(empty layer, schema ready)\nOther Nagaad cities\nOther hazards (e.g. conflict)")
    _card(s, 8.9, 1.15, 4.0, 2.3, "Not this product", "Probabilistic flood model\nClimate projections\nEngineering design flood\nInsurance / loss modelling")
    _bullets(
        s,
        [
            "GIS city name is Belet Weyne; MIS city name is Beledweyne — both are used on purpose",
            "Figures are a spatial overlay (intersects / length in flood), not a depth or return-period model",
        ],
        top=3.7,
        size=16,
    )

    # 5 KPIs
    s = _blank(prs)
    _title(s, "What users see: flood KPIs")
    _table(
        s,
        [
            ["KPI", "Beledweyne (21 Sep 2026)", "Meaning"],
            ["Flood area", "5,223 ha", "Historical flood mask inside the urban polygon"],
            ["Buildings in flood", "47,013 of 51,954", "Footprints intersecting the flood extent"],
            ["IDP people exposed", "97,813 (133 of 133 sites)", "People at sites that intersect flood"],
            ["OSM roads in flood", "410.3 km of 430.7 km", "City-wide mapped network length in flood"],
            ["Nagaad roads in flood", "22.7 km of 33.9 km", "Project completed / ongoing / design-ready"],
        ],
        0.4,
        1.2,
        12.5,
        3.4,
        col_w=[3.2, 4.0, 5.3],
    )
    _bullets(
        s,
        [
            "OSM and Nagaad kilometres are never mixed — two networks, two totals",
            "Priority table lists the 15 largest IDP sites in flood; click zooms the map",
        ],
        top=4.85,
        size=16,
    )

    # 6 Layers
    s = _blank(prs)
    _title(s, "Map layers")
    _table(
        s,
        [
            ["Group", "Layer", "How it is shown"],
            ["Flood hazard", "Historical flood extent", "Blue fill — Shabelle events clipped to the city"],
            ["Exposure", "Buildings", "Red in flood, grey outside (PMTiles)"],
            ["Exposure", "IDP and host sites", "Magenta circles sized by population"],
            ["Infrastructure", "OSM road network", "Brown in flood, grey outside"],
            ["Infrastructure", "Shabelle river + boundary", "Blue centreline; dashed urban extent"],
            ["Investment", "Nagaad project roads", "Green = completed/ongoing; purple = design ready"],
            ["Investment", "Nagaad project sites", "PIU office / lab"],
            ["Services (future)", "Hospitals, schools, markets", "Off until points are collected"],
        ],
        0.35,
        1.15,
        12.6,
        5.2,
        col_w=[2.4, 3.5, 6.7],
    )

    # 7 OSM vs Nagaad
    s = _blank(prs)
    _title(s, "Two road networks, analysed separately")
    _card(
        s,
        0.4,
        1.2,
        6.1,
        4.6,
        "OSM / city road network",
        "City-wide OpenStreetMap (SURP II roads clip)\n~4,888 segments · 430.7 km\nFlood join on the full network\nDelivered as compact PMTiles for the browser\nKPI: OSM roads in flood",
    )
    _card(
        s,
        6.8,
        1.2,
        6.1,
        4.6,
        "Nagaad project roads",
        "Only SURP-II / Nagaad alignments\nCompleted + ongoing GeoPackages\nDesign-ready KML / KMZ\n17 features · 33.9 km\nGreen vs purple by status\nKPI: Nagaad roads in flood",
        fill=PptRgb(0xEC, 0xFC, 0xFF),
    )

    # 8 Architecture
    s = _blank(prs)
    _title(s, "How it is put together")
    _card(s, 0.4, 1.2, 4.0, 2.5, "1. Raw GIS (local)", "~5.9 GB Beledweyne folder\nRead-only in Docker\nShapefiles, GPKG, KML/KMZ\nNever committed to Git")
    _card(s, 4.65, 1.2, 4.0, 2.5, "2. Docker ETL", "Python 3.12 + GeoPandas\nClip to Belet Weyne polygon\nFlood overlay + lengths\nGeoJSON, GPKG, PMTiles")
    _card(s, 8.9, 1.2, 4.0, 2.5, "3. Web dashboard", "Vite + React + TypeScript\nMapLibre GL + Esri basemap\nEmbed ?embed=1 in MIS\nGitHub Pages when published")
    _bullets(
        s,
        [
            "Converted reusable files: converted/beledweyne/  (GeoJSON + QGIS GeoPackage)",
            "Dashboard copies: public/data/beledweyne/  (small GeoJSON + PMTiles)",
            "MIS does not rebuild the map — it catalogues the URL and shows an iframe",
        ],
        top=4.0,
        size=16,
    )

    # 9 Tech
    s = _blank(prs)
    _title(s, "Technical stack")
    _table(
        s,
        [
            ["Layer", "Choice", "Why"],
            ["UI", "Vite 8, React 19, TypeScript, Tailwind 4", "Fast local app, typed, simple to embed"],
            ["Map", "MapLibre GL JS 6 + PMTiles", "No Mapbox token; large layers stay tiled"],
            ["Basemap", "Esri World Light Gray", "No API key; readable under overlays"],
            ["ETL", "Docker, GeoPandas, pyogrio, Shapely", "Reproducible GIS on any Windows/Mac host"],
            ["Tiles", "tippecanoe (felt)", "Buildings ~5.5 MB; OSM roads ~0.9 MB"],
            ["Host", "GitHub Pages + Actions", "Free, public, iframe-friendly"],
            ["MIS link", "Firestore riskDashboards + IframeContainer", "Access control stays in Nagaad MIS"],
        ],
        0.35,
        1.15,
        12.6,
        5.2,
        col_w=[2.0, 5.2, 5.4],
    )

    # 10 Pipeline
    s = _blank(prs)
    _title(s, "ETL pipeline (npm run etl)")
    _bullets(
        s,
        [
            "Bind the local raw folder read-only: D:\\Municipal Risk Dashboards\\DATA\\risk_dashboard\\beledweyne_data",
            "Select the SURP II city polygon UrbanName = Belet Weyne",
            "Clip historical flood, Shabelle river, buildings, OSM roads, IDP sites to that polygon",
            "Flag features that intersect flood; measure flooded road length in UTM 38N",
            "Convert SURP-II GeoPackages and Design Ready KML/KMZ to one Nagaad schema",
            "Write summary.json KPIs; PMTiles for flood (z0–14), buildings (z12–16), OSM roads (z10–16)",
            "Keep Nagaad roads as GeoJSON (small); drop huge intermediate building GeoJSON",
        ],
        size=16,
    )

    # 11 Conversion
    s = _blank(prs)
    _title(s, "Converting mixed investment formats")
    _table(
        s,
        [
            ["Source", "Format", "Converted to"],
            ["SURP-II completed / ongoing", "GeoPackage (already WGS84)", "GeoJSON + combined GPKG"],
            ["Design Ready Investments", "KML / KMZ (Google Earth)", "Same road schema, WGS84 GeoJSON"],
            ["OSM / SURP city roads", "Shapefile or GeoJSON", "osm_roads.geojson + roads.pmtiles"],
            ["Buildings", "Large shapefile", "buildings.pmtiles only"],
            ["Flood mask", "Shapefile", "Clipped / simplified PMTiles"],
            ["Hospitals / schools / markets", "Not collected yet", "Empty facilities.geojson + schema"],
        ],
        0.35,
        1.15,
        12.6,
        4.6,
        col_w=[3.6, 4.2, 4.8],
    )
    _bullets(
        s,
        ["Geometry CRS EPSG:4326. Lengths and areas EPSG:32638 (UTM 38N).", "Per-file GeoJSON kept under converted/beledweyne/sources/ for correction."],
        top=5.9,
        size=15,
    )

    # 12 Embed
    s = _blank(prs)
    _title(s, "How it sits in Nagaad MIS")
    _bullets(
        s,
        [
            "The dashboard is a standalone app — it is not rebuilt inside nagaad-mis-portal",
            "Local check first: http://localhost:5173/beledweyne?embed=1 (or the current Vite port)",
            "When published: https://nagaadproject.github.io/nagaad-municipal-risk/beledweyne?embed=1",
            "MIS Admin → Municipal Risk registers that URL for city Beledweyne",
            "The MIS iframe loads embed mode (compact header, no duplicate Nagaad chrome)",
            "GitHub Pages is built by Actions on push to main (workflow_dispatch also allowed)",
        ],
        size=17,
    )

    # 13 From screening to decisions
    s = _blank(prs)
    _title(s, "From overlay screening to municipal decisions")
    _table(
        s,
        [
            ["Stage", "What the city can do today", "What is still missing"],
            [
                "1. See",
                "Where flood, people, buildings, OSM and Nagaad roads overlap",
                "Depth, duration, neighbourhood names",
            ],
            [
                "2. Rank",
                "Largest IDP sites in flood; Nagaad km in flood vs OSM",
                "A transparent score (people + access + facilities + works)",
            ],
            [
                "3. Choose works",
                "See which Nagaad roads are ongoing vs design-ready",
                "Cost, drainage design, residual risk after the work",
            ],
            [
                "4. Prepare",
                "Know almost the whole urban area can flood",
                "Safe high ground, evacuation routes, early warning",
            ],
            [
                "5. Report",
                "Live map in MIS for discussion",
                "One-click briefing pack (map + table) for council",
            ],
        ],
        0.3,
        1.1,
        12.7,
        5.3,
        col_w=[1.8, 5.5, 5.4],
    )

    # 14 Product improvements
    s = _blank(prs)
    _title(s, "Dashboard improvements to request next")
    _bullets(
        s,
        [
            "Neighbourhood / district summaries (not only city-wide KPIs) so mayors can act by area",
            "Priority score combining IDP people, buildings, critical facilities, and Nagaad road km in flood",
            "Scenario flood layers (e.g. 2018 vs 2020 vs “deep water”) instead of one ever-flooded mask",
            "Critical facilities on the map: hospitals, health posts, schools, markets, water points, municipal offices",
            "Drainage, culverts, bridges, and river banks — the assets that actually fail in Beledweyne floods",
            "Export a one-page PDF/PNG for municipal meetings; Somali labels on the main UI",
            "Link each Nagaad road to MIS works (package, cost, contractor, status) so risk sits next to contracts",
            "Keep OSM and Nagaad separate in every chart and export — that split is already the right model",
        ],
        size=16,
    )

    # 15 Data request
    s = _blank(prs)
    _title(s, "Data to ask the analysis team for")
    _table(
        s,
        [
            ["Ask for", "Format", "Why municipalities need it"],
            ["Flood depth / return period (10–100 yr) or event years", "Raster or polygons", "Move from “wet/dry” to “how bad”"],
            ["Neighbourhoods / sections of Beledweyne", "Polygon + name", "Allocate machines, sandbags, and meetings by area"],
            ["Hospitals, schools, universities, markets, water", "Points: name, type, operator", "Protect services, not only buildings"],
            ["Drains, culverts, bridges, pumps, flood walls", "Lines/points + condition", "Find blockage and failure points"],
            ["DEM / low-lying pockets", "Raster or contours", "Where water stays after the river drops"],
            ["Nagaad works attributes", "Table joined to roads", "Cost, package, design standard, drainage Y/N"],
            ["Past flood damage / displacement", "Survey or points by year", "Ground the map in what actually happened"],
        ],
        0.28,
        1.1,
        12.75,
        5.4,
        col_w=[4.4, 2.6, 5.75],
    )

    # 16 Ready next
    s = _blank(prs)
    _title(s, "Ready next, and what this is not")
    _card(
        s,
        0.4,
        1.2,
        6.1,
        4.4,
        "Ready when data arrives",
        "Hospitals, schools/universities, markets\nSame flood join as other points\nMore Nagaad cities on the same pattern\nOptional later: conflict as a second hazard",
    )
    _card(
        s,
        6.8,
        1.2,
        6.1,
        4.4,
        "Disclaimer",
        "Indicative screening from a historical flood mask\nAlmost the whole urban area is in that mask\nNot design flood, depth, or annual probability\nUse to screen and discuss, not to replace engineering study",
        fill=PptRgb(0xFE, 0xF3, 0xC7),
    )

    prs.save(path)


def _set_cell_shading(cell, hex_color: str) -> None:
    tc = cell._tePr if False else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _shade_header(row) -> None:
    for cell in row.cells:
        _set_cell_shading(cell, "1E4D7B")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    _shade_header(table.rows[0])
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = val
    doc.add_paragraph()


def _heading(doc: Document, text: str, level=1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = DOC_NAVY


def build_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    header = section.header.paragraphs[0]
    header.text = ""
    if LOGO.exists():
        run = header.add_run()
        run.add_picture(str(LOGO), width=Inches(0.42))
        header.add_run("   Nagaad  ·  Somalia Urban Resilience Project Phase II  ·  Municipal Flood Risk Dashboard")
    else:
        header.text = "Nagaad  ·  Municipal Flood Risk Dashboard  ·  Process documentation"
    if header.runs:
        header.runs[-1].font.size = Pt(9)
        header.runs[-1].font.color.rgb = DOC_NAVY

    footer = section.footer.paragraphs[0]
    footer.text = f"Beledweyne flood screening  ·  {TODAY}  ·  Indicative, not a probabilistic model"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    if LOGO.exists():
        logo_p = doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = logo_p.add_run()
        run.add_picture(str(LOGO), width=Inches(1.25))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Nagaad Municipal Risk Dashboard")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = DOC_NAVY

    sub = doc.add_paragraph()
    run = sub.add_run("Process documentation — Beledweyne flood screening (this version)")
    run.font.size = Pt(14)
    run.font.color.rgb = DOC_NAVY

    meta = doc.add_paragraph()
    run = meta.add_run(
        f"Document date: {TODAY}\n"
        "Audience: Nagaad project team, municipal counterparts, and technical maintainers\n"
        "Companion briefing: Nagaad_Municipal_Risk_Dashboard_Briefing.pptx"
    )
    run.font.size = Pt(11)

    _heading(doc, "1. Purpose")
    doc.add_paragraph(
        "The Municipal Risk Dashboard is a city-scoped screening tool. For Beledweyne it answers: "
        "where does the historical Shabelle flood extent overlap buildings, IDP and host sites, the "
        "OpenStreetMap (OSM) road network, and Nagaad / SURP-II roads and sites?"
    )
    doc.add_paragraph(
        "It is a standalone open-source web application. Nagaad MIS does not redraw the map. The MIS "
        "catalogue stores the dashboard URL for city Beledweyne and displays it in an iframe "
        "(/portal/risk?city=Beledweyne). This version is flood-only. Conflict and other hazards are "
        "not processed or shown."
    )
    doc.add_paragraph(
        "Figures are indicative spatial overlays. They are not a design flood, a depth grid, or a "
        "probabilistic (return-period) risk model. Almost the entire Beledweyne urban polygon sits "
        "inside the dissolved historical flood mask, so many exposure shares are very high. That is "
        "a property of the mask, not a software error."
    )

    _heading(doc, "2. Naming")
    doc.add_paragraph(
        "The SURP II city polygon attribute UrbanName is “Belet Weyne”. The Nagaad MIS city label is "
        "“Beledweyne”. The ETL selects Belet Weyne for clipping and writes Beledweyne on the dashboard "
        "and in summary.json. Both names must stay in use."
    )

    _heading(doc, "3. What the application does")
    doc.add_paragraph("For a logged-in MIS user (or anyone opening the local/Pages URL) the app:")
    for item in [
        "Loads a light-grey Esri basemap (no API key) and the Beledweyne urban extent.",
        "Draws the historical flood mask clipped to that extent.",
        "Overlays buildings (in/out of flood), IDP sites, OSM roads, the Shabelle centreline, Nagaad project roads, and the PIU office.",
        "Shows five KPIs in one row: flood area, buildings in flood, IDP people exposed, OSM roads in flood, Nagaad roads in flood.",
        "Lists the 15 largest IDP sites in flood; a click flies the map to the site.",
        "Colours Nagaad roads by status: green = completed/ongoing, purple = design ready.",
        "Keeps an empty, switched-off layer for hospitals, schools/universities, and markets.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    _heading(doc, "4. Current Beledweyne snapshot")
    doc.add_paragraph(
        "Processed 21 September 2026 (UTC 09:30). Re-run ETL to refresh."
    )
    _add_table(
        doc,
        ["Indicator", "Value"],
        [
            ["Flood area inside city", "5,223 ha"],
            ["Buildings total / in flood", "51,954 / 47,013"],
            ["IDP sites total / in flood", "133 / 133"],
            ["IDP individuals in flood", "97,813"],
            ["IDP households in flood", "16,302"],
            ["OSM roads total / in flood", "430.7 km / 410.3 km"],
            ["Nagaad roads total / in flood", "33.9 km / 22.7 km"],
            ["Nagaad road features / intersecting flood", "17 / 17"],
            ["Nagaad sites (PIU office)", "1, outside flood mask"],
            ["OSM road features", "4,888"],
        ],
    )

    _heading(doc, "5. System architecture")
    doc.add_paragraph(
        "Three stores of data are kept strictly apart."
    )
    _add_table(
        doc,
        ["Store", "Location", "Role"],
        [
            [
                "Raw GIS",
                r"D:\Municipal Risk Dashboards\DATA\risk_dashboard\beledweyne_data",
                "Original shapefiles, GPKG, KML/KMZ (~5.9 GB). Read-only in Docker. Not in Git.",
            ],
            [
                "Converted geofiles",
                "converted/beledweyne/",
                "Reusable WGS84 GeoJSON, QGIS GeoPackage, conversion-log.json. For correction and reuse.",
            ],
            [
                "Dashboard payload",
                "public/data/beledweyne/",
                "Small GeoJSON + PMTiles + summary.json that the browser actually loads.",
            ],
        ],
    )
    doc.add_paragraph(
        "The web app is Vite 8 + React 19 + TypeScript + Tailwind 4. The map is MapLibre GL JS 6 with "
        "the PMTiles protocol (pmtiles 4). Basemap tiles are Esri World Light Gray Canvas. Large layers "
        "(flood, buildings, OSM roads) are vector tiles; small layers (boundary, river, IDPs, Nagaad "
        "roads and sites, empty facilities) are GeoJSON."
    )
    doc.add_paragraph(
        "ETL runs in Docker (python:3.12-bookworm, GeoPandas/pyogrio/Shapely, tippecanoe from felt). "
        "docker-compose bind-mounts the raw folder read-only, writes converted files to converted/beledweyne, "
        "writes dashboard files to public/data/beledweyne, and mounts etl/process.py so script edits do not "
        "require an image rebuild."
    )

    _heading(doc, "6. Software versions used for conversion")
    _add_table(
        doc,
        ["Tool", "Version / note"],
        [
            ["GeoPandas", "1.1.4"],
            ["pyogrio (GDAL/OGR)", "0.13.0"],
            ["pandas", "3.0.6"],
            ["Shapely", "2.x via GeoPandas"],
            ["tippecanoe", "felt/tippecanoe in the ETL image"],
            ["Coordinate system (geometry)", "EPSG:4326 WGS84"],
            ["Coordinate system (length/area)", "EPSG:32638 UTM zone 38N"],
        ],
    )

    _heading(doc, "7. End-to-end process")
    _heading(doc, "7.1 Prerequisites", 2)
    for item in [
        "Docker Desktop running.",
        "Raw Beledweyne GIS folder present on disk.",
        "Repository checkout with Node.js 22+ for the web app.",
        "Copy .env.example to .env. RAW_DATA_HOST must point at the raw folder using forward slashes, e.g. D:/Municipal Risk Dashboards/DATA/risk_dashboard/beledweyne_data.",
    ]:
        doc.add_paragraph(item, style="List Number")

    _heading(doc, "7.2 Run the ETL", 2)
    doc.add_paragraph("From the repository root:")
    p = doc.add_paragraph()
    run = p.add_run("npm run etl")
    run.bold = True
    doc.add_paragraph(
        "This is docker compose --env-file .env run --rm etl. The container entrypoint is python /etl/process.py."
    )

    _heading(doc, "7.3 City clip", 2)
    doc.add_paragraph(
        "Read raw_data/SURPII_city_polygons/SURPII_city_polygons.shp. Find column UrbanName (or urbanName). "
        "Keep the feature whose trimmed value equals Belet Weyne. Make the geometry valid and dissolve to a "
        "single city polygon. Write public/data/beledweyne/boundary.geojson labelled Beledweyne."
    )

    _heading(doc, "7.4 Flood mask", 2)
    doc.add_paragraph(
        "Read raw_data/som_floods/flood_historical/flood_extent_historical_dissolved.shp (historical flood "
        "extent). Reproject to WGS84 if needed, make valid, clip to the city polygon, simplify with tolerance "
        "0.00015 degrees (preserve topology), dissolve to one polygon, set inFlood = 1. Compute area in UTM 38N "
        "and store floodAreaHa. Build flood.pmtiles with tippecanoe minzoom 0, maxzoom 14, no feature/tile size "
        "limits (the mask is one polygon and must remain visible when zoomed in). Delete the intermediate "
        "flood.geojson after tiles are built."
    )

    _heading(doc, "7.5 River", 2)
    doc.add_paragraph(
        "Read raw_data/som_rivers/SOM_Rivers_Juba_shabelle.shp (source filename includes Juba; for Beledweyne "
        "only the Shabelle centreline remains after clip). Clip to the city. Keep name/code/class if present. "
        "Write river.geojson. The dashboard labels this layer “Shabelle river” only."
    )

    _heading(doc, "7.6 IDP and host sites", 2)
    doc.add_paragraph(
        "Prefer decisionMaking_data/belet_weyne_idps_flood_exposure.geojson if it exists; otherwise clip "
        "raw_data/som_idp/SURPII_city_idps.shp to UrbanName Belet Weyne. Spatial join to the flood polygon "
        "(intersects). Write idps.geojson. Dashboard popups use settlementName, settlementClass, idpIndividuals, "
        "idpHouseholds, inFlood."
    )

    _heading(doc, "7.7 Buildings", 2)
    doc.add_paragraph(
        "Read raw_data/som_buildings/buildings_beledweyne/Buildingfootprint_Beledweyne.shp. Make geometries valid. "
        "Keep area (area_in_me / areaM2) and confidence. Spatial join to flood. Write a temporary buildings.geojson "
        "then tippecanoe minzoom 12, maxzoom 16, drop-densest-as-needed. Delete the GeoJSON so the repo does not "
        "carry tens of thousands of polygons as text. The map uses buildings.pmtiles (~5.5 MB)."
    )

    _heading(doc, "7.8 OSM road network (separate from Nagaad)", 2)
    doc.add_paragraph(
        "Prefer decisionMaking_data/belet_weyne_roads_flood_exposure.geojson; else clip raw_data/som_roads/"
        "SURPII_city_roads.shp. Compute length in UTM 38N. Flag inFlood by intersection. Measure flooded length with "
        "polygon overlay (intersection, keep line type) and sum kilometres. Write network=osm. Copy the full "
        "attribute GeoJSON to converted/beledweyne/osm_roads.geojson. Build roads.pmtiles (minzoom 10, maxzoom 16). "
        "Delete the dashboard intermediate GeoJSON after tiling."
    )
    doc.add_paragraph(
        "KPI fields: osmRoadsKmTotal, osmRoadsKmInFlood (also copied to legacy keys roadsKmTotal / roadsKmInFlood)."
    )

    _heading(doc, "7.9 Nagaad investments — completed and ongoing (GeoPackage)", 2)
    doc.add_paragraph(
        "Walk folder SURP-II Investments completed and ongoing for every .gpkg. Read all OGR layers. Inputs already "
        "use EPSG:4326. Drop Z. Map attributes:"
    )
    _add_table(
        doc,
        ["Source field", "Dashboard field"],
        [
            ["subproject_name", "name"],
            ["package_name", "packageName"],
            ["category", "category (Roads / Buildings)"],
            ["status", "status (Completed / Ongoing)"],
            ["length_km", "lengthKm if > 0, else UTM length"],
        ],
    )
    doc.add_paragraph("Files processed:")
    _add_table(
        doc,
        ["File", "Category", "Status", "Treatment"],
        [
            ["Construction of Beledweyne PIU office and Lab.gpkg", "Buildings", "Completed", "Project site; outline polygonized if closed"],
            ["Irrid Amin Road IB (3km).gpkg", "Roads", "Ongoing", "Nagaad project road, 3.0 km"],
            ["Inner Ring Road (3.25km).gpkg", "Roads", "Ongoing", "Nagaad project road, 3.25 km"],
            ["Sheikh Hassan Barsane Road (3km).gpkg", "Roads", "Ongoing", "Nagaad project road, 3.0 km"],
        ],
    )
    doc.add_paragraph(
        "stage is set to completed_ongoing. network is nagaad. Each file is also written under "
        "converted/beledweyne/sources/ as its own GeoJSON."
    )

    _heading(doc, "7.10 Nagaad investments — design ready (KML / KMZ)", 2)
    doc.add_paragraph(
        "Walk folder Design Ready Investments for .kml and .kmz. KMZ is a zip containing doc.kml. A file named "
        ".kmz.kml is already-extracted KML and is read as KML. Driver: OGR KML via GeoPandas."
    )
    doc.add_paragraph("Normalisation:")
    for item in [
        "Name → name. If Name is missing or generic (e.g. Propose), use the file stem (e.g. Shabelow Bridge & Connected Roads).",
        "Strip KML chrome: timestamp, tessellate, extrude, visibility, drawOrder, icon, altitudeMode.",
        "status = Design ready; stage = design_ready; network = nagaad.",
        "category = Bridge if name or filename contains “bridge”, otherwise Roads.",
        "Road #11 YOBSAN MARKET contains five LineStrings; they stay five features.",
        "Length from UTM 38N (KML has no length_km).",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Design-ready inputs include Liiqliqato and Shabelow bridges, Irrid Amin Package 1A, Farah Afi 8A/8B, "
        "Buundo Labaad–Airport Road #4, Siigaalow roads #6 / #13A / #13B, and Yobsan Market Road #11."
    )

    _heading(doc, "7.11 Flood analysis on Nagaad roads and sites", 2)
    doc.add_paragraph(
        "Project roads and OSM roads are never concatenated before the flood overlay. Each collection is joined "
        "separately. inFlood = 1 if the feature intersects the flood polygon. Flooded kilometres use overlay "
        "intersection (keep line geometry) then UTM length. Sites use intersection only (no length). The PIU office "
        "does not intersect the mask (projectSitesInFlood = 0)."
    )
    doc.add_paragraph(
        "On the map, Nagaad road colour is status, not flood: green completed/ongoing, purple design ready. Flood "
        "exposure remains in KPIs and popups."
    )

    _heading(doc, "7.12 Outputs written by ETL", 2)
    doc.add_paragraph("converted/beledweyne/")
    for item in [
        "sources/*.geojson — one file per original GPKG/KML/KMZ",
        "osm_roads.geojson — OSM network + inFlood",
        "project_roads.geojson — 17 Nagaad road features",
        "project_sites.geojson — PIU office",
        "facilities.geojson — empty template",
        "facilities.schema.json — allowed types hospital | school | university | market",
        "beledweyne_nagaad_assets.gpkg — QGIS layers project_roads and project_sites",
        "conversion-log.json — machine log of inputs, tools, CRS, steps",
        "summary.json — KPI snapshot",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph("public/data/beledweyne/ (dashboard):")
    for item in [
        "boundary.geojson, river.geojson, idps.geojson",
        "project_roads.geojson, project_sites.geojson, facilities.geojson",
        "flood.pmtiles, buildings.pmtiles, roads.pmtiles",
        "summary.json, facilities.schema.json",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Conflict GeoJSON is deleted if a previous run created it. This version does not clip ACLED/conflict."
    )

    _heading(doc, "8. How to correct a converted geometry")
    for item in [
        "Edit the per-source GeoJSON in converted/beledweyne/sources/, or fix the original GPKG/KML and re-run ETL.",
        "If you edited sources only, merge back into project_roads.geojson / project_sites.geojson, or re-run ETL from originals.",
        "Copy the combined files into public/data/beledweyne/ if you skip a full ETL.",
        "Refresh the local dashboard. Do not write converted files into the raw GIS folder (it is read-only in Docker).",
    ]:
        doc.add_paragraph(item, style="List Number")

    _heading(doc, "9. Future facilities (hospitals, schools, markets)")
    doc.add_paragraph(
        "facilities.geojson is an empty FeatureCollection. Add Point features with name and type "
        "(hospital | school | university | market). Optional: status, source, inFlood. The map layer is off by default. "
        "If facilities are collected as a real GIS file, point the ETL at that file so a full re-run does not overwrite "
        "hand-edits with an empty template. Existing non-empty facilities.geojson is preserved and re-flagged for flood."
    )

    _heading(doc, "10. Web application implementation")
    _heading(doc, "10.1 City configuration", 2)
    doc.add_paragraph(
        "src/cities/beledweyne.ts holds centre [45.204, 4.736], zoom 13, minZoom 11, and layer definitions "
        "(labels, default on/off, legend colours). src/cities/types.ts defines CitySummary including the split OSM / "
        "Nagaad kilometre fields."
    )
    _heading(doc, "10.2 Data loading", 2)
    doc.add_paragraph(
        "src/lib/data.ts builds URLs under data/{slug}/. loadSummary fetches summary.json. loadGeoJSON fetches IDP "
        "GeoJSON for the priority table. MapLibre loads other GeoJSON and PMTiles by URL. PMTiles uses Protocol({ metadata: true }) "
        "and tilev4. MapLibre is excluded from Vite optimizeDeps so the worker module resolves."
    )
    _heading(doc, "10.3 Map", 2)
    doc.add_paragraph(
        "MapView waits for map loaded/idle before adding overlays (MapLibre can fire load before listeners attach). "
        "GeoJSON layers are added first; PMTiles sources are added next; vector layers attach when those sources are loaded. "
        "Nagaad roads are drawn above OSM roads. Embed mode uses a compact header."
    )
    _heading(doc, "10.4 Local run", 2)
    doc.add_paragraph("npm install then npm run dev. Open /beledweyne or /beledweyne?embed=1. If port 5173 is busy, Vite picks the next free port.")

    _heading(doc, "11. Publishing (when requested)")
    doc.add_paragraph(
        "Do not push until the team asks. Intended remote: nagaadproject/nagaad-municipal-risk. GitHub Actions "
        ".github/workflows/pages.yml builds with GITHUB_PAGES=true (base /nagaad-municipal-risk/) and deploys Pages. "
        "The workflow permission needs the workflow scope on the GitHub token. Embed URL:"
    )
    p = doc.add_paragraph()
    run = p.add_run("https://nagaadproject.github.io/nagaad-municipal-risk/beledweyne?embed=1")
    run.italic = True
    doc.add_paragraph(
        "In Nagaad MIS Admin → Municipal Risk, register that URL with city Beledweyne. Update any old rogaaltech Pages URL."
    )

    _heading(doc, "12. KPI definitions")
    _add_table(
        doc,
        ["KPI", "Definition"],
        [
            ["Flood area", "Area of the clipped historical flood polygon, UTM 38N, hectares"],
            ["Buildings in flood", "Count of building footprints whose geometry intersects flood"],
            ["IDP people exposed", "Sum of idpIndividuals at sites with inFlood = 1"],
            ["OSM roads in flood", "Length of OSM geometry overlapping flood, kilometres"],
            ["Nagaad roads in flood", "Length of Nagaad project road geometry overlapping flood, kilometres"],
        ],
    )

    _heading(doc, "13. Limitations")
    for item in [
        "Historical “ever flooded” mask, not hazard intensity, depth, or annual exceedance probability.",
        "Dissolved flood covers almost all of the Beledweyne urban polygon, so most buildings, IDP sites, and OSM roads flag inFlood.",
        "Nagaad road features may all intersect flood even when only part of their length is flooded (count vs kilometre).",
        "Building footprints and OSM roads follow SURP II / OSM quality; they are not a field inventory.",
        "Design-ready KML names are sometimes incomplete; file names are used as fallback.",
        "Conflict is excluded in this version even though ACLED data exist in the raw folder.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    _heading(doc, "14. Repository map")
    _add_table(
        doc,
        ["Path", "Role"],
        [
            ["etl/process.py", "Conversion and flood overlay"],
            ["etl/Dockerfile", "GeoPandas + tippecanoe image"],
            ["docker-compose.yml", "Binds raw, converted, public data, process.py"],
            ["src/components/MapView.tsx", "MapLibre overlays and colours"],
            ["src/components/KpiStrip.tsx", "Five flood KPIs"],
            ["docs/conversions.md", "Short conversion notes in-repo"],
            ["docs/nagaad-logo.png", "Nagaad SURP II seal used in dashboard, PPTX, DOCX"],
            ["docs/Nagaad_Municipal_Risk_Dashboard_Briefing.pptx", "This briefing deck"],
            ["docs/Nagaad_Municipal_Risk_Dashboard_Process.docx", "This document"],
        ],
    )

    _heading(doc, "15. Making this useful for municipal decisions")
    doc.add_paragraph(
        "Today the dashboard is a screening overlay. Municipalities can already see that flood, people, buildings, "
        "and roads occupy the same space, and they can see which Nagaad alignments are works versus design-ready. "
        "To become a decision model it must help a mayor or municipal engineer answer: where do we act first, with "
        "what type of work, and what residual risk remains after Nagaad investments."
    )
    _heading(doc, "15.1 Product improvements (dashboard)", 2)
    _add_table(
        doc,
        ["Improvement", "Decision it unlocks"],
        [
            ["Neighbourhood / district KPI split", "Allocate response and meetings by section of Beledweyne, not only city-wide"],
            ["Transparent priority score", "Rank IDP sites, facilities, and road packages on one comparable list"],
            ["Multi-scenario flood (events or return periods)", "Distinguish “sometimes wet” from “always / deep water”"],
            ["Critical facilities layer", "Protect hospitals, schools, markets, and water points, not only footprints"],
            ["Drainage, culverts, bridges, river banks", "Target the assets that fail and block flow"],
            ["One-page PDF/PNG export + Somali labels", "Brief council and communities without a live laptop"],
            ["Join Nagaad roads to MIS works (cost, contractor, package)", "Put risk next to the contract the municipality already tracks"],
            ["Evacuation / high-ground / access when flooded", "Plan movement, not only exposure counts"],
        ],
    )

    _heading(doc, "16. Data request for the analysis team")
    doc.add_paragraph(
        "Ask GIS/analysis to deliver WGS84 layers (GeoPackage or GeoJSON preferred; KML acceptable) clipped or "
        "attributable to Beledweyne, with a short data dictionary. Priority order:"
    )
    _add_table(
        doc,
        ["Priority", "Layer", "Minimum fields", "Use in the dashboard"],
        [
            ["P1", "Neighbourhoods / municipal sections", "name, population if known", "Filter KPIs and priority lists by area"],
            ["P1", "Hospitals and health posts", "name, type, operator, beds/staff if known", "Service continuity in flood"],
            ["P1", "Schools and universities", "name, type, enrolment if known", "Protect education sites and shelters"],
            ["P1", "Markets", "name, periodic/daily", "Protect livelihoods and access"],
            ["P1", "Flood by event year or return period, ideally with depth", "year or RP, depth class", "Replace single ever-flooded flag"],
            ["P2", "Drains, culverts, bridges, pumps", "type, condition, last cleared", "Find failure and blockage points"],
            ["P2", "DEM or low-lying pockets", "elevation or “ponding” class", "Where water remains after the river falls"],
            ["P2", "Nagaad works table joined to alignments", "package, cost USD, status, drainage Y/N, design standard", "Rank investments, not only lines"],
            ["P2", "Water points / berkads", "name, functional Y/N", "WASH during flood"],
            ["P3", "Past damage and displacement", "year, households, location", "Validate the map against lived events"],
            ["P3", "Rainfall / river gauge series", "station, date, value", "Link early warning later"],
            ["P3", "Solid-waste dump sites", "name, status", "Drainage blockage risk"],
        ],
    )
    doc.add_paragraph(
        "Facilities schema already reserved in facilities.geojson: name, type (hospital | school | university | market), "
        "status, source, inFlood. New facility types (health_post, water_point, municipal_office, mosque) can be added "
        "to that enum when data arrive. Neighbourhood polygons should use local names used by the municipality."
    )

    _heading(doc, "17. How municipalities should use the current version")
    for item in [
        "Use OSM vs Nagaad road km in flood to discuss whether project alignments are in the same hazard as the rest of the city (they are) and what drainage the design must include.",
        "Use the IDP priority table to sequence site visits, drainage around camps, and humanitarian coordination — not as a relocation order.",
        "Use green vs purple Nagaad roads to separate works already in construction from design-ready packages still able to change.",
        "Treat the 5,223 ha mask as “this city floods widely,” not as a precise building-level verdict. Confirm on the ground before any demolition, compensation, or final design.",
        "Do not use this version for insurance, land expropriation, or engineering freeboard. Those need depth, frequency, and survey.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "End of process documentation. For slide-level explanation of purpose and architecture, use the companion PowerPoint."
    )
    run.italic = True

    doc.save(path)


def main() -> None:
    pptx = OUT_DIR / "Nagaad_Municipal_Risk_Dashboard_Briefing.pptx"
    docx = OUT_DIR / "Nagaad_Municipal_Risk_Dashboard_Process.docx"
    build_pptx(pptx)
    build_docx(docx)
    print(f"wrote {pptx}")
    print(f"wrote {docx}")


if __name__ == "__main__":
    main()
